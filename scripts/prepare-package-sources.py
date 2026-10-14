"""Read NJO workbooks without changing them. Write a private, review-only import.

Run with the bundled Python runtime and explicit source paths. Never place output
under public/, dist/ or the source repository. Stock quantities are not inferred.
"""
import argparse
import hashlib
import json
import re
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from openpyxl import load_workbook
from docx import Document


def cents(value):
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    return int((Decimal(str(value)) * 100).quantize(Decimal('1'), rounding=ROUND_HALF_UP))


def prepare(args):
    products, packages, standards = {}, [], []
    paths = [Path(getattr(args, key)).resolve() for key in ('lifts', 'electrical', 'flatbeds', 'conventions')]
    sources = [{'file': p.name, 'sha256': hashlib.sha256(p.read_bytes()).hexdigest()} for p in paths]
    def add(key, title, category, source, sku=None, cost=None, retail=None, **extra):
        if key not in products:
            products[key] = {'sourceKey': key, 'title': title, 'category': category,
                'data': {'source': source['file'], 'sourceReferences': [source], 'supplierSku': str(sku) if sku is not None else None,
                    'candidateCostCents': cents(cost), 'candidateRetailCents': cents(retail), 'stockType': 'UNREVIEWED', 'unit': 'each', **extra}}
        elif source not in products[key]['data']['sourceReferences']:
            products[key]['data']['sourceReferences'].append(source)
            data = products[key]['data']
            for option in extra.get('options', []):
                if option not in data.setdefault('options', []):
                    data['options'].append(option)
            if data.get('candidateCostCents') != cents(cost) or data.get('candidateRetailCents') != cents(retail):
                data.setdefault('candidatePriceConflicts', []).append({'source': source, 'costCents': cents(cost), 'retailCents': cents(retail)})
        return key
    lift = load_workbook(paths[0], data_only=True, read_only=True)
    formula = load_workbook(paths[0], data_only=False, read_only=True)
    for sheet, title, audience, description in [
        ('Weekender', 'Weekender', 'Daily driving & weekends', 'A front suspension lift with refreshed front and rear shocks.'),
        ('Overlander (light)', 'Overlander Light', 'Gear goes on for trips', 'A complete suspension starting point for a truck that runs lighter between trips.'),
        ('Overlander (Heavy)', 'Overlander Heavy', 'Equipment stays onboard', 'A suspension starting point to specify around your permanent equipment and load.')]:
        s = lift[sheet]
        row = next(r for r in s.iter_rows(values_only=True) if ('Tacoma' in r and '16–23' in r))
        row_number = next(i for i, r in enumerate(s.iter_rows(values_only=True), 1) if r == row)
        c = lambda col: s[f'{col}{row_number}'].value
        ref = {'file': paths[0].name, 'sheet': sheet, 'row': row_number}
        total_formula = formula[sheet][f'{"E" if sheet == "Weekender" else "A"}{row_number}'].value
        rate_match = re.search(r'\*(\d+(?:\.\d+)?)', total_formula or '')
        if not rate_match:
            raise ValueError('Could not locate the source labor rate.')
        rate_cents = cents(float(rate_match.group(1)))
        lines, issues = [], []
        def line(role, name, sku_col, qty, cost_col, retail_col, supplier, key=None):
            sku = c(sku_col) if sku_col else None
            source = {**ref, 'range': f'{sku_col or cost_col}{row_number}:{retail_col}{row_number}'}
            product_key = add(key or f'lift:{supplier}:{sku}', name, 'Lifts & suspension', source, sku, c(cost_col), c(retail_col), supplier=supplier)
            lines.append({'sourceKey': product_key, 'role': role, 'quantity': qty, 'unitRetailCents': cents(c(retail_col))})
        if sheet == 'Weekender':
            line('Front shocks', 'B8 5100 front shock', 'N', c('O'), 'P', 'Q', 'Meyer Distributing')
            line('Rear shocks', 'B8 5100 rear shock', 'S', c('T'), 'U', 'V', 'Meyer Distributing')
            line('Strut mounts', 'Tacoma strut mount', 'W', c('X'), 'Y', 'Z', "O'Reilly")
            line('Lower mounting hardware', 'Tacoma OEM lower shock mounting bolt — VIN required', None, c('AD'), 'AE', 'AF', 'Toyota OEM', key='lift:toyota:vin-lower-shock-bolt')
            line('Differential drop', 'Differential drop kit', 'AJ', 1, 'AK', 'AL', 'Wheel Pros')
            hours, alignment = c('AN') + c('AP'), cents(c('AM'))
            listed = cents(c('G'))
            issues = [{'id': 'alignment-total', 'message': 'The source total counts alignment twice. Confirm the itemized price with one alignment charge.', 'resolved': False},
                      {'id': 'vin-hardware', 'message': 'The OEM lower mounting bolts require the customer VIN before ordering.', 'resolved': False}]
        else:
            line('Front suspension', 'B8 6112 front suspension kit', 'G', None, 'H', 'I', 'Meyer Distributing')
            line('Rear shocks', 'B8 5160 rear shock', 'K', c('L'), 'M', 'N', 'Meyer Distributing')
            line('Upper control arms', 'Upper control arm set', 'O', None, 'P', 'Q', 'Meyer Distributing')
            line('Brake lines', 'Extended brake lines', 'R', None, 'S', 'T', None)
            line('Differential drop', 'Differential drop kit', 'U', None, 'V', 'W', 'Wheel Pros')
            line('Cam bolt kit', 'Tacoma cam bolt kit', 'X', c('Y'), 'Z', 'AA', "O'Reilly")
            hours, alignment, listed = c('AF' if 'light' in sheet else 'AB'), None, None
            issues = [{'id': 'broken-total', 'message': 'The source total contains #REF! errors. Confirm missing quantities, labor scope and alignment before publishing a price.', 'resolved': False},
                      {'id': 'copied-description', 'message': 'The source description names 5100 shocks, but this row lists 6112 / 5160 parts. Confirm the customer-facing contents.', 'resolved': False}]
            if 'Heavy' in sheet:
                issues.append({'id': 'load-specification', 'message': 'The Light and Heavy rows list identical parts. Specify the permanent-load spring setup and confirm the difference before release.', 'resolved': False})
            else:
                issues.append({'id': 'variable-load', 'message': 'The package introduction promises variable-load rear springs, but this row does not specify them. Confirm the rear spring setup.', 'resolved': False})
        packages.append({'sourceKey': f'tacoma-2016-2023:{title.lower().replace(" ", "-")}', 'title': title, 'category': 'Lifts & suspension',
            'data': {'menuOrder': len(packages), 'audience': audience, 'description': description, 'vehicle': {'make': 'Toyota', 'model': 'Tacoma', 'yearFrom': 2016, 'yearTo': 2023},
                'lines': lines, 'labor': {'hours': hours, 'rateCents': rate_cents, 'alignmentCents': alignment}, 'source': ref,
                'sourceTotalCents': listed, 'issues': issues, 'fitmentNotes': '', 'installNotes': 'Confirm the VIN, factory suspension, tire clearance and loaded vehicle before ordering. Rust-related work is assessed separately.',
                'excludes': ['Tax', 'Freight', 'Tires and wheels', 'Additional rust-related work'], 'approvals': {'parts': False, 'pricing': False, 'fitment': False}}})
    flatbed = load_workbook(paths[2], data_only=True, read_only=True).active
    section = ''
    size_tags = {'Full-Size Long': 'FSL', 'Full-Size Regular': 'FSR', 'Mid-Size Long': 'MSL', 'Mid-Size Regular': 'MSR'}
    for number, row in enumerate(flatbed.iter_rows(values_only=True), 1):
        if number == 1:
            continue
        if row[0] and not row[1]:
            section = str(row[0]).strip()
        if not row[1] or not row[3]:
            continue
        sku = str(row[3]).strip()
        add('flatbed:' + sku, str(row[1]).strip(), 'Flatbeds', {'file': paths[2].name, 'sheet': flatbed.title, 'row': number},
            sku, row[13], row[14], candidateInstalledCents=cents(row[15]), options=[section],
            tags=['Flatbeds'] + ([size_tags[section]] if section in size_tags else []),
            description=str(row[7] or ''), changes=str(row[10] or ''), dimensions=str(row[5] or ''), weightKg=row[6], supplier=None)
    electrical = load_workbook(paths[1], data_only=True, read_only=True).active
    for number, row in enumerate(electrical.iter_rows(values_only=True), 1):
        if number < 4 or not row[0]:
            continue
        standards.append({'sourceKey': f'electrical-medium:{number}', 'title': str(row[0]), 'category': 'Power & electrical',
            'data': {'model': str(row[1] or ''), 'connections': str(row[2] or ''), 'purpose': str(row[3] or ''), 'bom': str(row[4] or ''),
                'source': {'file': paths[1].name, 'sheet': electrical.title, 'row': number}, 'scope': 'Medium electrical package; chassis battery run up to 20 ft.',
                'reviewNote': 'Source standard awaiting technical review. Confirm model numbers, terminal sizes, protection and wiring before installation.'}})
    conventions = [p.text.strip() for p in Document(paths[3]).paragraphs if p.text.strip()]
    standards.append({'sourceKey': 'catalog:shopify-conventions', 'title': 'Product types and sales options', 'category': 'Catalog',
        'data': {'model': 'Next Jump catalog conventions', 'scope': 'Preserve these labels in the native system. Tags alone do not approve prices or fitment.',
            'connections': 'Product types: Stocked Inventory, Non-Stocked Item, Service.\nSales options: Installed (Parts & Labor), DIY (Part Only).',
            'bom': 'Tags: Install-Only, Sellable, in-store-only, Flatbeds, FSL, FSR, MSL, MSR.\nVehicle collection tags are source labels and need exact fitment review.',
            'source': {'file': paths[3].name, 'sheet': 'Product types, tags and variants', 'row': None}}})
    return {'format': 'NJO_PACKAGE_WORKSHOP_V1', 'sourceId': hashlib.sha256(json.dumps(sources, sort_keys=True).encode()).hexdigest(),
        'sources': sources, 'products': list(products.values()), 'packages': packages, 'standards': standards,
        'conventions': {'source': paths[3].name, 'stockTypes': ['Stocked Inventory', 'Non-Stocked Item', 'Service'],
            'fulfillment': ['Installed (Parts & Labor)', 'DIY (Part Only)'], 'tags': ['Install-Only', 'Sellable', 'in-store-only', 'Flatbeds', 'FSL', 'FSR', 'MSL', 'MSR'],
            'vehicleTags': conventions[conventions.index('Vehicles for Flatbed Collections') + 1:conventions.index('Variants')]}}


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    for field in ('lifts', 'electrical', 'flatbeds', 'conventions', 'output'):
        parser.add_argument('--' + field, required=True)
    args = parser.parse_args()
    output = Path(args.output).resolve()
    root = Path(__file__).resolve().parents[1]
    if output == root or root in output.parents:
        raise SystemExit('Write the private import outside the website repository.')
    value = prepare(args)
    output.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
    output.write_text(json.dumps(value, ensure_ascii=False, indent=2))
    output.chmod(0o600)
    print(json.dumps({'parts': len(value['products']), 'packages': len(value['packages']), 'standards': len(value['standards']), 'state': 'DRAFT_ONLY', 'stockMovements': 0}))
